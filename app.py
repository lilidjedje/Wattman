import os
import sys
import importlib.util
import re
import unicodedata
from flask import Flask, render_template, request, redirect, url_for, jsonify, send_from_directory, send_file
from werkzeug.middleware.dispatcher import DispatcherMiddleware
from werkzeug.utils import secure_filename
import json
import shutil

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from io import BytesIO

app = Flask(__name__)

# Configuration
APP_DATA_DIR = os.environ.get('ENGINEERING_APP_DATA', '').strip()
UPLOAD_FOLDER = (
    os.path.join(APP_DATA_DIR, 'uploads')
    if APP_DATA_DIR
    else os.path.join(os.path.dirname(__file__), 'uploads')
)
ALLOWED_EXTENSIONS = {'xlsx', 'xls'}
MAX_CONTENT_LENGTH = 50 * 1024 * 1024  # 50MB max

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['CAN_SERVER_URL'] = os.environ.get('CAN_SERVER_URL', 'http://10.42.0.1/').strip() or 'http://10.42.0.1/'
app.config['MAX_CONTENT_LENGTH'] = MAX_CONTENT_LENGTH
app.config['MAX_FORM_MEMORY_SIZE'] = MAX_CONTENT_LENGTH
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

CAN_SERVER_DIR = os.path.join(os.path.dirname(__file__), 'can_server')
sys.path.insert(0, CAN_SERVER_DIR)
can_server_spec = importlib.util.spec_from_file_location('can_server_app', os.path.join(CAN_SERVER_DIR, 'app.py'))
can_server_module = importlib.util.module_from_spec(can_server_spec)
can_server_spec.loader.exec_module(can_server_module)
app.wsgi_app = DispatcherMiddleware(app.wsgi_app, {'/analyse_logs': can_server_module.app})

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def read_excel_file(filepath):
    """Lire le fichier Excel et retourner les données structurées"""
    try:
        workbook = load_workbook(filepath)
        sheets_data = {}
        
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            data = []
            headers = []
            
            # Obtenir le nombre maximal de colonnes utilisées
            max_col = worksheet.max_column
            max_row = worksheet.max_row
            
            # Obtenir les en-têtes de la première ligne
            first_row = worksheet[1]
            for idx in range(max_col):
                cell = first_row[idx] if idx < len(first_row) else None
                if cell and cell.value:
                    headers.append(str(cell.value))
                else:
                    headers.append(f"Colonne {idx + 1}")
            
            # Obtenir toutes les données (à partir de la ligne 2)
            for row_idx in range(2, max_row + 1):
                row_data = {}
                row = worksheet[row_idx]
                
                for col_idx in range(max_col):
                    cell = row[col_idx] if col_idx < len(row) else None
                    header = headers[col_idx]
                    
                    if cell:
                        value = cell.value
                    else:
                        value = None
                    
                    row_data[header] = value
                
                # Ajouter TOUTES les lignes, même si elles sont vides
                data.append(row_data)
            
            sheets_data[sheet_name] = {
                'headers': headers,
                'data': data,
                'row_count': len(data),
                'column_count': max_col
            }
        
        workbook.close()
        return sheets_data
    except Exception as e:
        return {'error': str(e)}

@app.route('/')
def index():
    """Page d'accueil"""
    return render_template('index.html')

@app.route('/import_project', methods=['GET'])
def import_project():
    """Page d'importation du projet"""
    return render_template('import_project.html')

@app.route('/select_components')
def select_components():
    """Sélection des Devices et ECU inclus dans les analyses du projet"""
    return render_template('select_components.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    """Traiter l'upload du fichier Excel"""
    if 'file' not in request.files:
        return jsonify({'error': 'Aucun fichier trouvé'}), 400
    
    file = request.files['file']
    
    if file.filename == '':
        return jsonify({'error': 'Aucun fichier sélectionné'}), 400
    
    if not allowed_file(file.filename):
        return jsonify({'error': 'Le fichier doit être un Excel (.xlsx ou .xls)'}), 400
    
    try:
        filename = secure_filename(file.filename)
        project_id = secure_filename(request.form.get('project_id', ''))
        target_folder = os.path.join(app.config['UPLOAD_FOLDER'], 'projects', project_id) if project_id else app.config['UPLOAD_FOLDER']
        os.makedirs(target_folder, exist_ok=True)
        filepath = os.path.join(target_folder, filename)
        file.save(filepath)
        
        # Lire les données du fichier
        sheets_data = read_excel_file(filepath)
        
        if 'error' in sheets_data:
            return jsonify({'error': sheets_data['error']}), 400
        
        # Retourner les données
        return jsonify({
            'success': True,
            'filename': filename,
            'projectId': project_id,
            'serverFileUrl': url_for('project_file', project_id=project_id, filename=filename) if project_id else url_for('uploaded_file', filename=filename),
            'sheets': sheets_data
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/uploads/<filename>')
def uploaded_file(filename):
    """Télécharger une fiche Wattman historique non encore associée à un identifiant projet."""
    return send_from_directory(app.config['UPLOAD_FOLDER'], secure_filename(filename), as_attachment=True)

@app.route('/project-files/<project_id>/<filename>')
def project_file(project_id, filename):
    """Télécharger la fiche Wattman conservée côté serveur pour un projet."""
    folder = os.path.join(app.config['UPLOAD_FOLDER'], 'projects', secure_filename(project_id))
    return send_from_directory(folder, secure_filename(filename), as_attachment=True)

@app.route('/recover-project-source', methods=['POST'])
def recover_project_source():
    payload = request.get_json(silent=True) or {}
    filename = secure_filename(str(payload.get('filename') or ''))
    project_id = secure_filename(str(payload.get('projectId') or ''))
    if not filename or not allowed_file(filename):
        return jsonify({'error': 'Nom de fiche Excel invalide.'}), 400
    candidates = []
    if project_id:
        candidates.append(os.path.join(app.config['UPLOAD_FOLDER'], 'projects', project_id, filename))
    candidates.append(os.path.join(app.config['UPLOAD_FOLDER'], filename))
    candidates.append(os.path.join(os.path.dirname(__file__), 'uploads', filename))
    for candidate in candidates:
        if not os.path.isfile(candidate):
            continue
        sheets = read_excel_file(candidate)
        if 'error' in sheets:
            return jsonify({'error': sheets['error']}), 400
        project_source = project_id and candidate == candidates[0]
        if not project_source:
            destination = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            if os.path.abspath(candidate) != os.path.abspath(destination):
                shutil.copy2(candidate, destination)
        source_url = url_for('project_file', project_id=project_id, filename=filename) if project_source else url_for('uploaded_file', filename=filename)
        return jsonify({'sheets': sheets, 'serverFileUrl': source_url})
    return jsonify({'error': 'La fiche Excel de cet ancien projet est introuvable. Ses sauvegardes ont été conservées.'}), 404


@app.route('/edit_current')
def edit_current():
    """Page d'édition du courant de homologation"""
    return render_template('edit_current.html')

@app.route('/homologation_configuration')
def homologation_configuration():
    """Configuration du courant de sortie DCDC utilisé par les synthèses Homologation."""
    return render_template('homologation_configuration.html')

@app.route('/save_current', methods=['POST'])
def save_current():
    """Sauvegarder les courants modifiés"""
    try:
        data = request.get_json()
        modifications = data.get('modifications', {})
        
        # Les modifications sont gérées côté client via sessionStorage
        return jsonify({
            'success': True,
            'message': 'Données sauvegardées'
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/display')
def display_excel():
    """Page d'affichage des données Excel"""
    return render_template('display_excel.html')

@app.route('/edit_alternator')
def edit_alternator():
    """Page d'édition des profils de l'alternateur"""
    return render_template('edit_alternator.html')

@app.route('/edit_pao')
def edit_pao():
    """Page de saisie du nombre et caractéristiques de PAO"""
    return render_template('edit_pao.html')

@app.route('/edit_gmv')
def edit_gmv():
    """Page de configuration du nombre et caractéristiques des GMV"""
    return render_template('edit_gmv.html')

@app.route('/mode_selection')
def mode_selection():
    """Page intermédiaire de sélection de l'outil d'analyse"""
    return render_template('mode_selection.html', can_server_url=app.config['CAN_SERVER_URL'])

@app.route('/synthesis')
def synthesis():
    """Page de rendu des tableaux de synthèse de calculs globaux"""
    return render_template('synthesis.html')

@app.route('/export_synthesis_excel', methods=['POST'])
def export_synthesis_excel():
    """Conserver le classeur Wattman et y injecter les synthèses Homologation recalculées."""
    try:
        payload = request.get_json(silent=True) or {}
        project_id = secure_filename(str(payload.get('projectId') or ''))
        filename = secure_filename(str(payload.get('filename') or ''))
        synthesis = payload.get('synthesis') or {}

        source_path = None
        if filename and project_id:
            candidate = os.path.join(app.config['UPLOAD_FOLDER'], 'projects', project_id, filename)
            if os.path.isfile(candidate):
                source_path = candidate
        if source_path is None and filename:
            candidate = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            if os.path.isfile(candidate):
                source_path = candidate

        workbook = load_workbook(source_path) if source_path else Workbook()
        if source_path is None and workbook.active:
            workbook.remove(workbook.active)

        thin = Side(style='thin', color='333333')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        section_fills = {
            'winter': '1BD719', 'midseason': '16AEDE', 'summer': 'E9DD19',
            'cycle1': '1BD719', 'cycle2': '16AEDE', 'cycle3': 'E9DD19'
        }
        titles = {
            'winter': 'Winter (<5 °C)', 'midseason': 'Mid-season (5 °C<T<15 °C)', 'summer': 'Summer (T>15 °C)',
            'cycle1': 'Cycle 1', 'cycle2': 'Cycle 2', 'cycle3': 'Cycle 3'
        }

        def write_synthesis_sheet(sheet_name, title, groups):
            if sheet_name in workbook.sheetnames:
                position = workbook.sheetnames.index(sheet_name)
                workbook.remove(workbook[sheet_name])
                worksheet = workbook.create_sheet(sheet_name, position)
            else:
                worksheet = workbook.create_sheet(sheet_name)

            worksheet.sheet_view.showGridLines = False
            worksheet.merge_cells('A1:D1')
            worksheet['A1'] = title
            worksheet['A1'].font = Font(name='Times New Roman', size=18, bold=True, color='FFFFFF')
            worksheet['A1'].fill = PatternFill('solid', fgColor='1A1D1F')
            worksheet['A1'].alignment = Alignment(horizontal='center', vertical='center')
            worksheet.row_dimensions[1].height = 28

            row_index = 3
            for group_key, records in groups.items():
                worksheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=4)
                group_cell = worksheet.cell(row=row_index, column=1, value=titles.get(group_key, group_key))
                group_cell.font = Font(name='Times New Roman', bold=True, color='111111')
                group_cell.fill = PatternFill('solid', fgColor=section_fills.get(group_key, 'E9DD19'))
                group_cell.alignment = Alignment(horizontal='center')
                row_index += 1

                headers = ['Time', 'Base Vehicle global', 'Upper Vehicle global', 'DCDC output']
                for column, header in enumerate(headers, 1):
                    cell = worksheet.cell(row=row_index, column=column, value=header)
                    cell.font = Font(name='Times New Roman', bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='4A4F52')
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.border = border
                row_index += 1

                for record in records:
                    values = [record.get('time', 0), record.get('base', 0), record.get('upper', 0), record.get('dcdcOutput', 0)]
                    for column, value in enumerate(values, 1):
                        cell = worksheet.cell(row=row_index, column=column, value=float(value))
                        cell.font = Font(name='Times New Roman', color='111111')
                        cell.fill = PatternFill('solid', fgColor='FFFFFF' if row_index % 2 else 'EDF4F7')
                        cell.alignment = Alignment(horizontal='center')
                        cell.border = border
                        cell.number_format = '0.## "min"' if column == 1 else '0.00 "A"'
                    row_index += 1
                row_index += 2

            for column, width in enumerate((16, 24, 24, 20), 1):
                worksheet.column_dimensions[get_column_letter(column)].width = width

        write_synthesis_sheet(
            'Synthèse DCDC Sizing',
            'Synthèse DCDC Sizing',
            synthesis.get('dcdc') or {}
        )
        write_synthesis_sheet(
            'Synthèse WLTC -7C',
            'Synthèse WLTC (-7 °C)',
            synthesis.get('wltc') or {}
        )

        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name='Synthese_Homologation.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as exc:
        return jsonify({'error': str(exc)}), 500

@app.route('/export_customer_results_excel', methods=['POST'])
def export_customer_results_excel():
    """Exporter les quatre feuilles Customer visibles du classeur Consoveille."""
    workbook = None
    try:
        source_path = os.path.join(os.path.dirname(__file__), 'Consoveille.xlsx')
        if not os.path.isfile(source_path):
            return jsonify({'error': 'Le classeur source Consoveille.xlsx est introuvable.'}), 404

        workbook = load_workbook(source_path, data_only=False, keep_links=True)
        customer_mode_index = workbook.sheetnames.index('Customer Mode')
        if customer_mode_index == 0:
            return jsonify({'error': "Aucune feuille ne précède « Customer Mode »."}), 500

        visible_sheets = [
            workbook.sheetnames[customer_mode_index - 1],
            'Customer Mode',
            'Customer wo unplung',
            'Customer PIE'
        ]
        missing_sheets = [name for name in visible_sheets if name not in workbook.sheetnames]
        if missing_sheets:
            return jsonify({
                'error': f"Feuille(s) introuvable(s) dans Consoveille.xlsx : {', '.join(missing_sheets)}"
            }), 500

        # Les feuilles dépendantes restent dans le classeur, mais sont masquées afin
        # de préserver les formules et de ne présenter que les quatre onglets demandés.
        for worksheet in workbook.worksheets:
            worksheet.sheet_state = 'visible' if worksheet.title in visible_sheets else 'veryHidden'

        workbook.active = workbook.sheetnames.index(visible_sheets[0])
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = 'auto'

        output = BytesIO()
        workbook.save(output)
        workbook.close()
        workbook = None
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name='Consoveille_Resultats_Customer.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except ValueError:
        return jsonify({'error': 'La feuille « Customer Mode » est introuvable dans Consoveille.xlsx.'}), 500
    except Exception as exc:
        return jsonify({'error': str(exc)}), 500
    finally:
        if workbook is not None:
            workbook.close()

@app.route('/export_dm_results_excel', methods=['POST'])
def export_dm_results_excel():
    """Exporter les résultats DM dans les pages Logistic Wo unplug et Delivery de Consoveille."""
    workbook = None
    try:
        payload = request.get_json(silent=True) or {}
        results = payload.get('results') or []
        mode = str(payload.get('mode') or 'DM1')
        project = payload.get('project') or {}
        source_path = os.path.join(os.path.dirname(__file__), 'Consoveille.xlsx')
        if not os.path.isfile(source_path):
            return jsonify({'error': 'Le classeur source Consoveille.xlsx est introuvable.'}), 404

        workbook = load_workbook(source_path, data_only=False, keep_links=False)
        visible_sheets = ['Logistic Wo unplug', 'Delivery']
        missing = [name for name in visible_sheets if name not in workbook.sheetnames]
        if missing:
            return jsonify({'error': f"Feuille(s) introuvable(s) : {', '.join(missing)}"}), 500

        for worksheet in workbook.worksheets:
            worksheet.sheet_state = 'visible' if worksheet.title in visible_sheets else 'veryHidden'
        workbook.active = workbook.sheetnames.index('Logistic Wo unplug')

        def number(value):
            try:
                return float(value)
            except (TypeError, ValueError):
                return 0.0

        metrics = {
            'AUTONOMIE': [number(item.get('autonomy')) for item in results],
            'CAPACITE BATTERIE': [number(item.get('capacity')) for item in results],
            'SOC DE DEPART': [number(item.get('startSoc')) for item in results],
            'SOC MINIMUM': [number(item.get('minSoc')) for item in results],
            'ENERGIE DISPONIBLE BRUTE': [number(item.get('gross')) for item in results],
            'ENERGIE PAR ENDORMISSEMENT': [number(item.get('vehicleUnit')) for item in results],
            'CONSO VEILLE INSTANTANEE': [number(item.get('instantaneous')) for item in results],
            'CONSOMMATION COURANT VEILLE': [number(item.get('sleepCurrentEnergy')) for item in results],
            'REVEIL TECHNIQUE': [number(item.get('technicalWake')) for item in results],
            'CONSO ALARME': [number(item.get('alarmCurrent')) for item in results],
            'CONSOMMATION ALARME': [number(item.get('alarmEnergy')) for item in results],
            'SANS E BDU': [number(item.get('sansEbduCurrent')) for item in results],
            'CONSOMMATION SANS E BDU': [number(item.get('sansEbduEnergy')) for item in results],
            'TWT CONSO IVC': [number(item.get('twtIvc')) for item in results],
            'POWERLATCH': [number(item.get('twtPowerlatch')) for item in results],
            'DUREE D ENVOI': [number(item.get('twtDuration')) for item in results],
            'CONSOMMATION EN VEILLE TOTALE': [number(item.get('dailyTotal')) for item in results]
        }

        def normalize(value):
            plain = unicodedata.normalize('NFKD', str(value or ''))
            plain = ''.join(char for char in plain if not unicodedata.combining(char))
            return re.sub(r'[^A-Z0-9]+', ' ', plain.upper()).strip()

        for worksheet in (workbook['Logistic Wo unplug'], workbook['Delivery']):
            # Métadonnées du projet, uniquement quand les libellés existent dans le modèle.
            for row in worksheet.iter_rows():
                for cell in row:
                    label = normalize(cell.value)
                    if label in {'PROJET', 'PROJECT'}:
                        worksheet.cell(cell.row, min(cell.column + 1, worksheet.max_column), str(project.get('name') or ''))
                    elif label in {'MODE', 'MODE AUTONOMIE'}:
                        worksheet.cell(cell.row, min(cell.column + 1, worksheet.max_column), mode)

            for row in worksheet.iter_rows():
                row_label = ' '.join(normalize(cell.value) for cell in row if isinstance(cell.value, str))
                if not row_label:
                    continue
                matching = next((values for fragment, values in metrics.items() if fragment in row_label), None)
                if not matching:
                    continue
                candidates = [
                    cell for cell in row
                    if cell.__class__.__name__ != 'MergedCell'
                    and (isinstance(cell.value, (int, float)) or (isinstance(cell.value, str) and cell.value.startswith('=')))
                ]
                for index, value in enumerate(matching[:len(candidates)]):
                    candidates[index].value = value

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = 'auto'
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        workbook = None
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=f'Consoveille_Resultats_{secure_filename(mode)}.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as exc:
        return jsonify({'error': str(exc)}), 500
    finally:
        if workbook is not None:
            workbook.close()

@app.route('/export_dcdc_template_excel', methods=['POST'])
def export_dcdc_template_excel():
    """Injecter les résultats Homologation dans les quatre pages du modèle DCDC."""
    workbook = None
    try:
        payload = request.get_json(silent=True) or {}
        synthesis = payload.get('synthesis') or {}
        devices = payload.get('devices') or []
        project = payload.get('project') or {}
        template_path = os.path.join(os.path.dirname(__file__), 'DCDC.xlsx')
        if not os.path.isfile(template_path):
            return jsonify({'error': 'Le modèle DCDC.xlsx est introuvable.'}), 404

        workbook = load_workbook(template_path, data_only=False, keep_links=True)
        if len(workbook.worksheets) < 6:
            return jsonify({'error': 'Le modèle DCDC.xlsx ne contient pas les six onglets attendus.'}), 500

        update_sheet, definition_sheet, dcdc_sheet, dcdc_graph, wltc_sheet, wltc_graph = workbook.worksheets[:6]
        previous_wltc_graph_title = wltc_graph.title
        wltc_sheet.title = 'Bilan -7°C'
        wltc_graph.title = 'Graph -7°C'
        for chart in wltc_graph._charts:
            for series in chart.ser:
                references = [
                    getattr(getattr(series, 'val', None), 'numRef', None),
                    getattr(getattr(series, 'cat', None), 'numRef', None),
                    getattr(getattr(series, 'tx', None), 'strRef', None)
                ]
                for reference in references:
                    if reference is not None and reference.f:
                        reference.f = reference.f.replace(previous_wltc_graph_title, wltc_graph.title)
        update_sheet.sheet_state = 'veryHidden'
        definition_sheet.sheet_state = 'veryHidden'
        for worksheet in (dcdc_sheet, dcdc_graph, wltc_sheet, wltc_graph):
            worksheet.sheet_state = 'visible'
        workbook.active = workbook.index(dcdc_sheet)

        def number(value, fallback=0):
            try:
                return float(value)
            except (TypeError, ValueError):
                return fallback

        def normalized(value):
            plain = unicodedata.normalize('NFKD', str(value or ''))
            plain = ''.join(character for character in plain if not unicodedata.combining(character))
            return re.sub(r'[^A-Z0-9]+', '', plain.upper())

        def write_project_metadata(worksheet):
            worksheet['C2'] = str(project.get('name') or '')
            worksheet['F3'] = str(project.get('date') or '')
            worksheet['F4'] = str(project.get('pilotWattman') or '')

        write_project_metadata(dcdc_sheet)
        wltc_sheet['C3'] = str(project.get('name') or '')
        wltc_sheet['C4'] = str(project.get('date') or '')
        wltc_sheet['C5'] = str(project.get('pilotWattman') or '')

        dcdc_profile_columns = {
            'winter': list(range(8, 13)),
            'midseason': list(range(14, 19)),
            'summer': list(range(20, 25))
        }
        wltc_profile_columns = {
            'cycle1': list(range(10, 17)),
            'cycle2': list(range(18, 24)),
            'cycle3': list(range(25, 31))
        }

        valid_devices = [
            device for device in devices
            if str(device.get('name') or '').strip()
        ]

        def synchronize_equipment_rows(
            worksheet, start_row, end_row, profile_columns, clear_columns, write_device
        ):
            all_profile_columns = [
                column for columns in profile_columns.values() for column in columns
            ]
            existing_device_rows = []
            empty_rows = []
            for row_index in range(start_row, end_row + 1):
                if any(
                    worksheet.cell(row_index, column).__class__.__name__ == 'MergedCell'
                    for column in clear_columns
                ):
                    continue
                name_cell = worksheet.cell(row_index, 2)
                name = name_cell.value
                has_profile = any(
                    worksheet.cell(row_index, column).value is not None
                    for column in all_profile_columns
                )
                if name and has_profile:
                    existing_device_rows.append(row_index)
                elif not name and name_cell.__class__.__name__ != 'MergedCell':
                    empty_rows.append(row_index)

            pending_by_name = {}
            for device in valid_devices:
                pending_by_name.setdefault(normalized(device.get('name')), []).append(device)

            assigned = {}
            available_rows = []
            for row_index in existing_device_rows:
                key = normalized(worksheet.cell(row_index, 2).value)
                candidates = pending_by_name.get(key) or []
                if candidates:
                    assigned[row_index] = candidates.pop(0)
                else:
                    available_rows.append(row_index)

            assigned_device_ids = {id(device) for device in assigned.values()}
            remaining_devices = [
                device
                for device in valid_devices
                if id(device) not in assigned_device_ids
            ]
            available_rows.extend(
                row_index for row_index in empty_rows
                if row_index not in available_rows
            )
            if len(remaining_devices) > len(available_rows):
                raise ValueError(
                    f"La feuille « {worksheet.title} » ne contient pas assez de lignes "
                    "pour tous les Devices, PAO et GMV."
                )
            for row_index, device in zip(available_rows, remaining_devices):
                assigned[row_index] = device

            # La liste exportée est autoritaire : toute ancienne ligne d'équipement
            # absente de l'outil est vidée avant de réécrire la sélection courante.
            managed_rows = set(existing_device_rows) | set(assigned)
            for row_index in managed_rows:
                for column in clear_columns:
                    worksheet.cell(row_index, column).value = None
            for row_index, device in sorted(assigned.items()):
                write_device(row_index, device)

        def write_dcdc_device(row_index, device):
            name = str(device.get('name') or '').strip()
            current = number(device.get('current'))
            lower = int(number(device.get('lower'), 1))
            upper = int(number(device.get('upper'), 0))
            dcdc_sheet.cell(row_index, 2, name)
            dcdc_sheet.cell(row_index, 5, lower)
            dcdc_sheet.cell(row_index, 6, upper)
            dcdc_sheet.cell(row_index, 7, current)
            for profile, columns in dcdc_profile_columns.items():
                values = device.get('dcdcProfiles', {}).get(profile, [])
                numeric_values = [number(value) for value in values]
                for column, value in zip(columns, numeric_values):
                    dcdc_sheet.cell(row_index, column, value)
                if numeric_values:
                    dcdc_sheet.cell(row_index, columns[-1] + 1, max(numeric_values))

        def write_wltc_device(row_index, device):
            name = str(device.get('name') or '').strip()
            current = number(device.get('current'))
            lower = int(number(device.get('lower'), 1))
            upper = int(number(device.get('upper'), 0))
            wltc_sheet.cell(row_index, 2, name)
            wltc_sheet.cell(row_index, 5, lower)
            wltc_sheet.cell(row_index, 6, upper)
            wltc_sheet.cell(row_index, 7, lower)
            wltc_sheet.cell(row_index, 8, upper)
            wltc_sheet.cell(row_index, 9, current)
            for profile, columns in wltc_profile_columns.items():
                values = device.get('wltcProfiles', {}).get(profile, [])
                for column, value in zip(columns, values):
                    wltc_sheet.cell(row_index, column, number(value))

        synchronize_equipment_rows(
            dcdc_sheet, 
            9,
            84,
            dcdc_profile_columns,
            [2, 5, 6, 7, *range(8, 26)],
            write_dcdc_device
        )
        synchronize_equipment_rows(
            wltc_sheet,
            15,
            176,
            wltc_profile_columns,
            [2, *range(5, 32)],
            write_wltc_device
        )

        def records(group, profile):
            values = group.get(profile) or []
            return [{
                'time': number(record.get('time')),
                'base': number(record.get('base')),
                'upper': number(record.get('upper')),
                'dcdcOutput': number(record.get('dcdcOutput'), 250)
            } for record in values]

        # Les cellules de synthèse du modèle alimentent le bilan DCDC.
        for profile, columns in (
            ('winter', list(range(8, 13))),
            ('midseason', list(range(14, 19))),
            ('summer', list(range(20, 25)))
        ):
            profile_records = records(synthesis.get('dcdc') or {}, profile)
            interval_records = profile_records[::2][:len(columns)]
            for column, record in zip(columns, interval_records):
                dcdc_sheet.cell(92, column, record['upper'])
                dcdc_sheet.cell(93, column, record['base'])
            if interval_records:
                maximum_column = columns[-1] + 1
                dcdc_sheet.cell(92, maximum_column, max(record['upper'] for record in interval_records))
                dcdc_sheet.cell(93, maximum_column, max(record['base'] for record in interval_records))

        dcdc_graph_layout = {
            'winter': (1, 2, 3, 4),
            'midseason': (6, 7, 8, 9),
            'summer': (11, 12, 13, 14)
        }
        for profile, columns in dcdc_graph_layout.items():
            profile_records = records(synthesis.get('dcdc') or {}, profile)
            for offset in range(12):
                row_index = 4 + offset
                if offset < len(profile_records):
                    record = profile_records[offset]
                    for column, key in zip(columns, ('time', 'base', 'upper', 'dcdcOutput')):
                        dcdc_graph.cell(row_index, column, record[key])
                else:
                    for column in columns:
                        dcdc_graph.cell(row_index, column).value = None

        # Les bilans WLTC et leurs trois graphiques utilisent les cycles 1 à 3.
        wltc_summary_columns = {
            'cycle1': list(range(10, 17)),
            'cycle2': list(range(18, 24)),
            'cycle3': list(range(25, 31))
        }
        for profile, columns in wltc_summary_columns.items():
            profile_records = records(synthesis.get('wltc') or {}, profile)
            interval_records = profile_records[::2][:len(columns)]
            for column, record in zip(columns, interval_records):
                wltc_sheet.cell(94, column, record['upper'])
                wltc_sheet.cell(95, column, record['base'])
            if interval_records:
                wltc_sheet.cell(94, 31, max(record['upper'] for record in interval_records))
                wltc_sheet.cell(95, 31, max(record['base'] for record in interval_records))

        wltc_graph_layout = {
            'cycle1': (1, 2, 3, 4),
            'cycle2': (6, 7, 8, 9),
            'cycle3': (11, 12, 13, 14)
        }
        wltc_graph['A9'] = 'Cycle 1'
        wltc_graph['F9'] = 'Cycle 2'
        wltc_graph['K9'] = 'Cycle 3'
        for profile, columns in wltc_graph_layout.items():
            profile_records = records(synthesis.get('wltc') or {}, profile)
            for offset in range(14):
                row_index = 11 + offset
                if offset < len(profile_records):
                    record = profile_records[offset]
                    for column, key in zip(columns, ('time', 'base', 'upper', 'dcdcOutput')):
                        wltc_graph.cell(row_index, column, record[key])
                else:
                    for column in columns:
                        wltc_graph.cell(row_index, column).value = None

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = 'auto'

        output = BytesIO()
        workbook.save(output)
        workbook.close()
        workbook = None
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name='Synthese_Homologation_DCDC.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as exc:
        return jsonify({'error': str(exc)}), 500
    finally:
        if workbook is not None:
            workbook.close()

@app.route('/export_synthesis_pdf')
def export_synthesis_pdf():
    """Prévisualisation et export PDF de toutes les synthèses du projet."""
    return render_template('export_synthesis_pdf.html')

@app.route('/battery_menu')
def battery_menu():
    """Page du sous-menu Autonomie Batterie"""
    return render_template('battery_menu.html')

@app.route('/edit_customer_mode')
def edit_customer_mode():
    """Page de configuration des profils exclusivement pour le Customer Mode"""
    return render_template('edit_customer_mode.html')

@app.route('/edit_dm_mode')
def edit_dm_mode():
    """Page dédiée de configuration des profils pour DM1 et DM2"""
    return render_template('edit_dm_mode.html')

@app.route('/biberonnage')
def biberonnage():
    """Choix de la durée ou tableau éditable correspondant."""
    days = request.args.get('days', type=int)
    if days not in (30, 90):
        return render_template('biberonnage_selection.html')
    return render_template('biberonnage.html', duration_days=days)

@app.route('/biberonnage/resultats')
def biberonnage_resultats():
    """Synthèse des Ah dépensés pour le tableau de biberonnage."""
    days = request.args.get('days', default=30, type=int)
    if days not in (30, 90):
        days = 30
    return render_template('biberonnage_resultats.html', duration_days=days)

@app.route('/battery_results_menu')
def battery_results_menu():
    """Page de choix du type d'analyse d'autonomie avec fenêtres modales"""
    return render_template('battery_results_menu.html')

@app.route('/dm_results')
def dm_results():
    """Synthèse dédiée aux profils d'autonomie DM1 et DM2."""
    return render_template('dm_results.html')

@app.route('/fixed_consumption')
def fixed_consumption():
    """Configuration des consommations forfaitaires Customer Mode"""
    return render_template('fixed_consumption.html')

@app.route('/sleep_consumption')
def sleep_consumption():
    """Configuration des consommations en veille Customer Mode"""
    return render_template('sleep_consumption.html')

@app.route('/battery_configuration')
def battery_configuration():
    """Page de configuration des caractéristiques de la batterie et des réveils"""
    return render_template('battery_configuration.html')

@app.route('/vehicle_sleep')
def vehicle_sleep():
    """Compatibilité des anciens liens : l'étape a été supprimée."""
    return redirect(url_for('battery_menu'))

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5001)
