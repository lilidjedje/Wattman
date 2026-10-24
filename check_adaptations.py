import sys
sys.path.insert(0, 'requirements')
sys.path.insert(0, 'requirements/openpyxl')
from openpyxl import load_workbook

wb = load_workbook('default-project.xlsx')

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} ===')
    print(f'Max row: {ws.max_row}, Max col: {ws.max_column}')
    
    # Chercher ADAPTATIONS dans cette feuille
    found = False
    for row in ws.iter_rows(min_row=1, values_only=False):
        for cell in row:
            if cell.value and 'ADAPT' in str(cell.value).upper():
                if not found:
                    found = True
                    print(f'Trouvé ADAPTATIONS:')
                print(f'  Ligne {cell.row}, Col {cell.column}: {cell.value}')
                # Afficher la ligne complète
                ligne = []
                for i, c in enumerate(ws[cell.row]):
                    ligne.append(c.value)
                print(f'  Valeurs ligne {cell.row}: {ligne}')
