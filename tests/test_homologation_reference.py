import json
import pathlib
import subprocess
import unittest

import openpyxl


ROOT = pathlib.Path(__file__).resolve().parents[1]


class HomologationReferenceTests(unittest.TestCase):
    def test_six_tables_match_sumproduct_with_large_currents(self):
        workbook = openpyxl.load_workbook(ROOT / 'DCDC - Test calcul.xlsx', data_only=True)
        sheet = workbook.worksheets[2]
        devices = []
        for row in range(9, 85):
            current = sheet.cell(row, 7).value
            if not isinstance(current, (int, float)) or 'GMV REFROIDISSEMENT' in str(sheet.cell(row, 2).value):
                continue
            devices.append(dict(key=str(row), rowIndex=row, name=sheet.cell(row, 2).value,
                                current=current, lower=sheet.cell(row, 5).value, upper=sheet.cell(row, 6).value))
        javascript = """
const fs = require('fs');
const vm = require('vm');
const devices = JSON.parse(fs.readFileSync(0, 'utf8'));
const storage = {homologationDevices: JSON.stringify(devices)};
const context = vm.createContext({window: {}, sessionStorage: {
    getItem: key => storage[key] ?? null, setItem: (key, value) => {storage[key] = value;}
}, document: {getElementById: () => ({appendChild() {}, getContext() {}}), createElement: () => ({}), querySelectorAll: () => []}, Chart: function () {}});
vm.runInContext(fs.readFileSync('static/homologation_defaults.js', 'utf8'), context);
context.HomologationDefaults = context.window.HomologationDefaults;
const html = fs.readFileSync('templates/synthesis.html', 'utf8');
vm.runInContext(html.slice(html.indexOf('        const devTimeline'), html.indexOf('        function buildExcelExportDevices')), context);
console.log(vm.runInContext('JSON.stringify(synthesisExportData)', context));
"""
        result = subprocess.run(['node', '-e', javascript], cwd=ROOT, input=json.dumps(devices),
                                text=True, capture_output=True, check=True)
        actual = json.loads(result.stdout)
        midseason_start = actual['dcdc']['midseason'][:2]
        self.assertEqual([record['time'] for record in midseason_start], [0, 5])
        for record in midseason_start:
            self.assertAlmostEqual(record['base'], 1077.4, places=3)
        self.assertEqual(midseason_start[0]['base'], midseason_start[1]['base'])
        self.assertEqual(midseason_start[0]['upper'], midseason_start[1]['upper'])
        for mode, sheet_index, current_column, row_start, row_end, profiles in [
            ('dcdc', 2, 7, 9, 84, dict(winter=range(8, 13), midseason=range(14, 19), summer=range(20, 25))),
            ('wltc', 4, 9, 15, 86, dict(cycle1=range(10, 17), cycle2=range(18, 24), cycle3=range(25, 31))),
        ]:
            sheet = workbook.worksheets[sheet_index]
            for profile, columns in profiles.items():
                for interval, column in enumerate(columns):
                    for field, active_column in [('base', 5), ('upper', 6)]:
                        expected = sum(sheet.cell(row, active_column).value * sheet.cell(row, current_column).value
                                       * (sheet.cell(row, column).value or 0)
                                       for row in range(row_start, row_end + 1)
                                       if isinstance(sheet.cell(row, active_column).value, (int, float))
                                       and isinstance(sheet.cell(row, current_column).value, (int, float)))
                        with self.subTest(mode=mode, profile=profile, interval=interval, field=field):
                            self.assertAlmostEqual(actual[mode][profile][interval * 2][field], expected, places=3)

    def test_default_profiles_match_reference_equipment(self):
        workbook = openpyxl.load_workbook(ROOT / 'DCDC.xlsx', data_only=True)
        requests = []
        for mode, sheet_index, start, end, columns in [
            ('dcdc', 2, 9, 84, [8, 9, 10, 11, 12, 14, 15, 16, 17, 18, 20, 21, 22, 23, 24]),
            ('wltc', 4, 15, 86, [10, 11, 12, 13, 14, 15, 16, 18, 19, 20, 21, 22, 23, 25, 26, 27, 28, 29, 30]),
        ]:
            sheet = workbook.worksheets[sheet_index]
            occurrences = {}
            for row in range(start, end + 1):
                name = sheet.cell(row, 2).value
                values = [sheet.cell(row, column).value for column in columns]
                if not name or not any(isinstance(value, (float, int)) for value in values):
                    continue
                if 'GMV REFROIDISSEMENT' in name:
                    continue
                occurrence = occurrences.get(name, 0)
                occurrences[name] = occurrence + 1
                requests.append(dict(mode=mode, name=name, occurrence=occurrence,
                                     expected=[round((value or 0) * 100, 8) for value in values]))
        javascript = """
const fs = require('fs');
global.window = {};
require('./static/homologation_defaults.js');
const requests = JSON.parse(fs.readFileSync(0, 'utf8'));
console.log(JSON.stringify(requests.map(request => window.HomologationDefaults.valuesFor(request.mode, request.name, request.occurrence))));
"""
        result = subprocess.run(['node', '-e', javascript], cwd=ROOT,
                                input=json.dumps(requests), text=True, capture_output=True, check=True)
        for request, actual in zip(requests, json.loads(result.stdout)):
            with self.subTest(mode=request['mode'], name=request['name']):
                self.assertEqual(actual, request['expected'])


if __name__ == '__main__':
    unittest.main()
