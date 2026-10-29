"""Refresh the read-only Excel reference profiles and independent test checkpoints."""
import json
import math
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

root = Path(__file__).resolve().parents[1]
workbook = openpyxl.load_workbook(root / 'Consoveille.xlsx', read_only=True, data_only=True)
title = next(name for name in workbook.sheetnames if name.startswith('Biberonnage_23'))
data = list(workbook[title].iter_rows(max_row=284, max_col=132, values_only=True))


def value(row, column):
    result = data[row - 1][column - 1]
    return result if isinstance(result, (int, float)) else 0


mapping = {1: 15, 2: 16, 3: 17, 4: 18, 5: 19, 6: 20, 7: 21, 8: 22, 9: 23, 10: 24, 11: 25,
           13: 26, 14: 27, 15: 28, 16: 29, 17: 30, 18: 31, 19: 32, 20: 33, 21: 34, 22: 35,
           24: 36, 25: 37, 26: 38, 27: 39, 28: 41}
subday = [('30s', '30 s'), ('1min', '1 min'), ('2min', '2 min'), ('5min', '5 min'),
          ('7min', '7 min'), ('1h', '1 h'), ('8h', '8 h')]
columns = []
for column in range(5, 106):
    day = value(11, column) or value(12, column)
    if not day or day > 90:
        continue
    phase_change = day >= 1 and abs(day - round(day)) > 1e-9
    if column <= 11:
        key, label = subday[column - 5]
    elif phase_change:
        key, label = f'day-{math.floor(day)}-after', f'{math.floor(day)} (après)'
    else:
        key, label = f'day-{round(day)}', str(round(day))
    columns.append(dict(key=key, label=label, daysElapsed=day,
                        day=math.floor(day) if day >= 1 else None,
                        dayNum=round(day) if day >= 1 and not phase_change else None,
                        sourceColumn=get_column_letter(column), phaseChange=phase_change))
reference = dict(source='Consoveille.xlsx', sheet=title, columns=columns,
                 currents={index: [value(row, col + 5) for col in range(len(columns))]
                           for index, row in mapping.items()},
                 deliveryTwtAh=[value(265, col + 5) for col in range(len(columns))])
fixtures = {}
for key, daily_row, soc_row, duration in [('base30', 213, 217, 30), ('tt30', 238, 242, 30), ('delivery90', 266, 270, 90)]:
    indices = [i for i, col in enumerate(columns) if col['daysElapsed'] <= duration and col['day']]
    fixtures[key] = dict(sheet=title, dailyRow=daily_row, socRow=soc_row,
                         checkpoints=[dict(key=columns[i]['key'], ah=value(daily_row, i + 5), soc=value(soc_row, i + 5))
                                      for i in indices if columns[i]['dayNum']],
                         sumAh=sum(value(daily_row, i + 5) for i in indices))
(root / 'static/biberonnage_reference.js').write_text(
    '// Generated from Consoveille.xlsx by tests/extract_biberonnage_reference.py; source inputs only.\n'
    'window.BiberonnageReference = ' + json.dumps(reference, ensure_ascii=False, separators=(',', ':')) + ';\n', encoding='utf-8')
(root / 'tests/fixtures/biberonnage_reference.json').write_text(json.dumps(fixtures, indent=2), encoding='utf-8')
workbook.close()
print(len(columns), {key: item['sumAh'] for key, item in fixtures.items()})
